from abc import ABC, abstractmethod
from indexer.schemas.document import RawDocument, ParsedDocument
from indexer.utils.logging import setup_logger
from io import BytesIO
from PyPDF2 import PdfReader
import docx
import json

class BaseParser(ABC):
    """Base class for all parser implementations"""

    def __init__(self):
        """Initialize the base parser."""
        self.logger = setup_logger(name=f"Indexer: {self.__class__.__name__.lower()}_parser")

    @abstractmethod
    def parse(self, raw_document):
        """Parse a raw document.

        Args:
            raw_document: Raw document to parse

        Returns:
            Parsed document
        """
        pass

class TextParser(BaseParser):
    """Parser for text documents."""

    def parse(self, raw_document: RawDocument) -> ParsedDocument:
        """Parse a raw text document.

        Args:
            raw_document: Raw text document to parse

        Returns:
            Parsed text document
        """
        try:
            self.logger.debug(f"Parsing text document: {raw_document.file_name}")
            try:
                parsed_content = raw_document.content.decode("utf-8")
            except UnicodeDecodeError:
                self.logger.warning(f"UTF-8 decode failed, falling back to latin-1 for {raw_document.file_name}")
                parsed_content = raw_document.content.decode("latin-1")
            except Exception as e:
                self.logger.error(f"Error decoding text document: {str(e)}")
                raise e

            self.logger.debug(f"Successfully parsed text document: {raw_document.file_name}")
            return ParsedDocument(
                user_id=raw_document.user_id,
                space_name=raw_document.space_name,
                file_name=raw_document.file_name,
                blob_path=raw_document.blob_path,
                content=parsed_content
                #content_type="text/plain"
            )
        except Exception as e:
            self.logger.error(f"Error parsing text document: {str(e)}")
            raise e

class PDFParser(BaseParser):
    """Parser for PDF documents."""

    def parse(self, raw_document: RawDocument) -> ParsedDocument:
        """Parse a raw PDF document.

        Args:
            raw_document: Raw PDF document to parse

        Returns:
            Parsed PDF document
        """
        try:
            self.logger.debug(f"Parsing PDF document: {raw_document.file_name}")
            content_stream = BytesIO(raw_document.content)
            pdf_reader = PdfReader(content_stream)
            
            self.logger.debug(f"PDF has {len(pdf_reader.pages)} pages")
            parsed_content = "\n".join(
                page.extract_text()
                for page in pdf_reader.pages
                if page.extract_text()
            )
            
            self.logger.debug(f"Successfully parsed PDF document: {raw_document.file_name}")
            return ParsedDocument(
                user_id=raw_document.user_id,
                space_name=raw_document.space_name,
                file_name=raw_document.file_name,
                blob_path=raw_document.blob_path,
                content=parsed_content
                #content_type="application/pdf"
            )
        except Exception as e:
            self.logger.error(f"Error parsing PDF document: {str(e)}")
            raise e

class DocxParser(BaseParser):
    """Parser for DOCX documents."""

    def parse(self, raw_document: RawDocument) -> ParsedDocument:
        """Parse a raw DOCX document.

        Args:
            raw_document: Raw DOCX document to parse

        Returns:
            Parsed DOCX document
        """
        try:
            self.logger.debug(f"Parsing DOCX document: {raw_document.file_name}")
            content_stream = BytesIO(raw_document.content)
            doc = docx.Document(content_stream)
            
            self.logger.debug(f"DOCX has {len(doc.paragraphs)} paragraphs")
            parsed_content = "\n".join(
                paragraph.text
                for paragraph in doc.paragraphs
                if paragraph.text
            )
            
            self.logger.debug(f"Successfully parsed DOCX document: {raw_document.file_name}")
            return ParsedDocument(
                user_id=raw_document.user_id,
                space_name=raw_document.space_name,
                file_name=raw_document.file_name,
                blob_path=raw_document.blob_path,
                content=parsed_content
                #content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
        except Exception as e:
            self.logger.error(f"Error parsing DOCX document: {str(e)}")
            raise e

class MarkdownParser(BaseParser):
    """Parser for Markdown documents."""

    def parse(self, raw_document: RawDocument) -> ParsedDocument:
        """Parse a raw Markdown document.

        Args:
            raw_document: Raw Markdown document to parse

        Returns:
            Parsed Markdown document
        """
        try:
            self.logger.debug(f"Parsing Markdown document: {raw_document.file_name}")
            try:
                parsed_content = raw_document.content.decode("utf-8")
            except UnicodeDecodeError:
                self.logger.warning(f"UTF-8 decode failed, falling back to latin-1 for {raw_document.file_name}")
                parsed_content = raw_document.content.decode("latin-1")
            except Exception as e:
                self.logger.error(f"Error decoding Markdown document: {str(e)}")
                raise e
            
            self.logger.debug(f"Successfully parsed Markdown document: {raw_document.file_name}")
            return ParsedDocument(
                user_id=raw_document.user_id,
                space_name=raw_document.space_name,
                file_name=raw_document.file_name,
                blob_path=raw_document.blob_path,
                content=parsed_content
                #content_type="text/markdown"
            )
        except Exception as e:
            self.logger.error(f"Error parsing Markdown document: {str(e)}")
            raise e

class JSONParser(BaseParser):
    """Parser for JSON documents."""

    def parse(self, raw_document: RawDocument) -> ParsedDocument:
        """Parse a raw JSON document.

        Args:
            raw_document: Raw JSON document to parse

        Returns:
            Parsed JSON document
        """
        try:
            self.logger.debug(f"Parsing JSON document: {raw_document.file_name}")
            parsed_content = json.loads(raw_document.content)
            
            self.logger.debug(f"Successfully parsed JSON document: {raw_document.file_name}")
            return ParsedDocument(
                user_id=raw_document.user_id,
                space_name=raw_document.space_name,
                file_name=raw_document.file_name,
                blob_path=raw_document.blob_path,
                content=parsed_content
                #content_type="application/json"
            )
        except Exception as e:
            self.logger.error(f"Error parsing JSON document: {str(e)}")
            raise e

# Define factory class
class ContentParser(BaseParser):
    """Factory for document parsers based on file extension."""

    def __init__(self):
        """Initialize the content parser factory."""
        super().__init__()
        self._parsers: dict[str, BaseParser] = {
            "txt": TextParser(),
            "pdf": PDFParser(),
            "docx": DocxParser(),
            "md": MarkdownParser(),
            "json": JSONParser()
        }
        self.logger = setup_logger(name="content_parser")
        self.logger.debug(f"Initialized content parser with parsers for: {', '.join(self._parsers.keys())}")

    def parse(self, raw_document: RawDocument) -> ParsedDocument:
        """Parse a raw document.

        Args:
            raw_document: Raw document to parse

        Returns:
            Parsed document
        """
        try:
            # check if parser exists for the given content type
            if raw_document.file_extension not in self._parsers:
                self.logger.error(f"No parser found for content type: {raw_document.file_extension}")
                raise ValueError(f"No parser found for content type: {raw_document.file_extension}\n, available parsers: {self._parsers.keys()}")
            
            self.logger.debug(f"Using {raw_document.file_extension} parser for {raw_document.file_name}")
            parser = self._parsers[raw_document.file_extension]
            return parser.parse(raw_document)
        except Exception as e:
            self.logger.error(f"Error parsing document content: {str(e)}")
            raise e
